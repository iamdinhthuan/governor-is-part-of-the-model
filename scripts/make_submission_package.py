"""Assemble the JSA submission package and the Overleaf upload folder.

  <pkg>/01_JSA_main.pdf                 pdfLaTeX build of the LaTeX source below
  <pkg>/02_figures_JSA/Figure_N.pdf     one vector file per figure, in paper order
  <pkg>/03_highlights_JSA.docx|.txt
  <pkg>/04_graphical_abstract_JSA.tif   (from make_graphical_abstract.py)
  <pkg>/05_cover_letter_JSA.pdf|.txt
  <pkg>/06_LaTeX_source_JSA.zip         JSA_Overleaf_Source/ incl. main_jsa.bbl
  <pkg>/07_declaration_of_interest_JSA.docx
  <pkg>/08_author_vitae_JSA_TEMPLATE.docx
  <pkg>/README_FIRST.txt, SHA256SUMS.txt
  <overleaf>/JSA_Overleaf_Source/ and JSA_Overleaf_Source.zip (same bytes as 06)

Run: uv run --with python-docx --with pymupdf --with pillow --with matplotlib \
       python make_submission_package.py <pkg_dir> <overleaf_dir>
TEXBIN must point at a TeX Live bin directory with pdflatex and bibtex.
"""

import datetime
import hashlib
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import docx
import pymupdf
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent
PAPER = next(d / "paper" for d in (ROOT, ROOT.parent) if (d / "paper" / "main.tex").exists())
PKG = Path(sys.argv[1]).resolve()
OVL = Path(sys.argv[2]).resolve()
TEXBIN = os.environ.get("TEXBIN", "/tmp/tinytex_test/TinyTeX/bin/universal-darwin")
DATE = "7 October 2026"
ZIP_TIME = (2026, 10, 7, 12, 0, 0)
ENV = dict(os.environ, PATH=TEXBIN + os.pathsep + os.environ["PATH"],
           SOURCE_DATE_EPOCH="1791345600", FORCE_SOURCE_DATE="1")
TITLE = ("The governor is part of the model: duty-cycle-aware deployment of "
         "object detectors on edge accelerators")
AUTHORS = [
    ("Dinh Thuan Nguyen", "Faculty of Electrical and Electronics Engineering, Ton Duc Thang University, Ho Chi Minh City, Vietnam"),
    ("Lam Phuong Nguyen", "Faculty of Electrical and Electronics Engineering, Ton Duc Thang University, Ho Chi Minh City, Vietnam"),
    ("Vinh Huy Nguyen", "Center of Visualization and Simulation, Duy Tan University, Da Nang, Vietnam"),
    ("Sy Vu Quang", "Advanced Intelligent Technology Research Group, Faculty of Electrical and Electronics Engineering, Ton Duc Thang University, Ho Chi Minh City, Vietnam"),
    ("Mohan Rajesh Elara", "ROAR Lab, Engineering Product Development, Singapore University of Technology and Design, Singapore"),
    ("Anh Vu Le", "Advanced Intelligent Technology Research Group, Faculty of Electrical and Electronics Engineering, Ton Duc Thang University, Ho Chi Minh City, Vietnam"),
]


def new_docx(title: str):
    d = docx.Document()
    c = d.core_properties
    c.author = c.last_modified_by = "Anh Vu Le"
    c.title, c.comments, c.subject = title, "", TITLE
    c.created = c.modified = datetime.datetime(*ZIP_TIME)
    c.revision = 1
    return d


def run(cmd, cwd):
    return subprocess.run(cmd, cwd=cwd, env=ENV, capture_output=True, text=True)


def latex_build(src: Path, name: str, bib: bool) -> Path:
    """pdflatex (+bibtex) until labels settle; fail on any unresolved item."""
    steps = [["pdflatex", "-interaction=nonstopmode", "-halt-on-error", name]]
    if bib:
        steps.append(["bibtex", name])
    steps += [["pdflatex", "-interaction=nonstopmode", "-halt-on-error", name]] * 3
    for cmd in steps:
        r = run(cmd, src)
        if r.returncode != 0:
            sys.exit(f"{cmd[0]} failed in {src}:\n{r.stdout[-3000:]}")
    log = (src / f"{name}.log").read_text(errors="replace")
    bad = re.findall(r"(?:Reference|Citation) `[^']*' on page \d+ undefined"
                     r"|Label\(s\) may have changed|There were multiply-defined labels"
                     r"|Missing character|^! .*", log, re.M)
    if bad:
        sys.exit(f"{name}.log: {bad[:5]}")
    return src / f"{name}.pdf"


def check_pdf(path: Path):
    d = pymupdf.open(path)
    t3 = {f[3] for p in d for f in p.get_fonts() if f[2] == "Type3"}
    if t3:
        sys.exit(f"{path.name}: Type 3 fonts {sorted(t3)}")
    return d


def strip_tex(s: str) -> str:
    s = s.replace("$2.4\\times$", "2.4\u00d7").replace("\\%", "%").replace("--", "\u2013")
    assert "\\" not in s and "$" not in s, s
    return s.strip()


def write_zip(src_dir: Path, zpath: Path):
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for p in [src_dir] + sorted(src_dir.rglob("*")):
            arc = p.relative_to(src_dir.parent).as_posix() + ("/" if p.is_dir() else "")
            zi = zipfile.ZipInfo(arc, ZIP_TIME)
            zi.external_attr = ((0o40755 if p.is_dir() else 0o100644) << 16)
            if p.is_dir():
                z.writestr(zi, b"")
            else:
                zi.compress_type = zipfile.ZIP_DEFLATED
                z.writestr(zi, p.read_bytes())


# ---------------------------------------------------------------------------
def build_source(tmp: Path) -> tuple[Path, Path]:
    subprocess.run([sys.executable, str(ROOT / "make_overleaf.py"), str(tmp)], check=True)
    src = tmp / "JSA_Overleaf_Source"
    build = tmp / "build"
    shutil.copytree(src, build)
    latex_build(build, "main_jsa", bib=True)
    shutil.copy2(build / "main_jsa.bbl", src / "main_jsa.bbl")
    return src, build


def figures(build: Path, out: Path):
    tex = (build / "main_jsa.tex").read_text()
    aux = (build / "main_jsa.aux").read_text()
    out.mkdir()
    rows = []
    for env in re.findall(r"\\begin\{figure\*?\}(.*?)\\end\{figure\*?\}", tex, re.S):
        img = re.search(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]+)\}", env).group(1)
        lab = re.search(r"\\label\{([^}]+)\}", env).group(1)
        num = re.search(r"\\newlabel\{" + re.escape(lab) + r"\}\{\{(\d+)\}", aux).group(1)
        shutil.copy2(build / "figures" / img, out / f"Figure_{num}.pdf")
        rows.append((int(num), img, lab))
    rows.sort()
    assert [r[0] for r in rows] == list(range(1, len(rows) + 1)), rows
    for f in out.iterdir():
        check_pdf(f)
    return rows


def highlights(out: Path):
    tex = (PAPER / "main.tex").read_text()
    body = re.search(r"\\begin\{highlights\}(.*?)\\end\{highlights\}", tex, re.S).group(1)
    items = [strip_tex(x) for x in body.split("\\item") if x.strip()]
    assert 3 <= len(items) <= 5 and all(len(x) <= 85 for x in items), items
    (out / "03_highlights_JSA.txt").write_text("".join(f"- {x}\n" for x in items))
    d = new_docx("Highlights")
    d.styles["Normal"].font.name = "Times New Roman"
    d.styles["Normal"].font.size = Pt(12)
    d.add_paragraph("Highlights")
    for x in items:
        d.add_paragraph(x, style="List Bullet")
    d.save(out / "03_highlights_JSA.docx")
    return items


LETTER = [
    "Dear Editors of the Journal of Systems Architecture,",
    "Please consider our manuscript, \u201c" + TITLE + ",\u201d for publication as a research "
    "article in the Journal of Systems Architecture: Embedded Software Design.",
    "Edge object detectors are usually compared by the latency of the compiled engine. The "
    "manuscript shows that on an embedded GPU this number does not predict the cost inside an "
    "application, because the stock frequency governors respond to the application\u2019s duty "
    "cycle. On a Jetson Orin Nano Super, the FP16 TensorRT engine of YOLO26n takes 4.2 ms "
    "back-to-back but 13.2 ms inside a sequential application, because the GPU governor holds "
    "the GPU at its minimum clock. Every measurement covers the full decode-to-detections "
    "pipeline, with module input power, clocks, GPU load and official COCO accuracy on all "
    "5,000 val2017 images.",
    "The paper makes three main points. First, a duty-cycle model describes the measurements: "
    "engine time scales with the inverse of the clock (3.7% mean absolute error), and a "
    "sequential pipeline released from the top clock falls to the floor within seconds. "
    "Second, restructuring only the host pipeline, under default governors and with "
    "byte-identical predictions, gives 2.4\u00d7 the throughput of the best sequential "
    "pipeline at locked maximum clocks. On the sequential and prefetch schedules, a sleeping "
    "CUDA wait lets the CPU governor lower the host clock and costs 17\u201320% of "
    "throughput; a per-task utilization clamp, set without root privileges, recovers most of "
    "this loss. Third, replications on a Raspberry Pi 5 with a Hailo-8 NPU and on a Qualcomm "
    "QCM6490 show that the pipeline gains transfer (3.3\u00d7 on the Pi 5), while the "
    "mechanism moves to the host.",
    "The work fits the scope of the Journal of Systems Architecture because it concerns the "
    "interaction between system software (frequency governors, the CPU scheduler and "
    "accelerator synchronization primitives) and the architecture of the application "
    "pipeline on embedded accelerators. It also gives a measurement method for reporting "
    "edge detection efficiency for the whole pipeline under deployment governors.",
    "The manuscript is original, is not under consideration elsewhere, and all six authors "
    "have approved its submission. The research received no specific grant from funding "
    "agencies in the public, commercial, or not-for-profit sectors. Measurement scripts, raw "
    "traces, per-run reports and the generators of every table and figure are available at "
    "https://github.com/iamdinhthuan/governor-is-part-of-the-model; the prediction files whose "
    "hashes anchor the byte-identical-prediction claims are archived on Zenodo "
    "(DOI 10.5281/zenodo.23185860). The manuscript includes data and code availability, "
    "CRediT, funding, competing-interest, and generative-AI declarations.",
    "Thank you for your consideration.",
]
SIGN = ["Sincerely,", "", "Anh Vu Le", "Corresponding author",
        "Advanced Intelligent Technology Research Group",
        "Faculty of Electrical and Electronics Engineering, Ton Duc Thang University",
        "Ho Chi Minh City, Vietnam", "leanhvu@tdtu.edu.vn"]


def to_latex(s: str) -> str:
    s = s.replace("%", "\\%").replace("\u00d7", "$\\times$").replace("\u2013", "--")
    s = s.replace("\u201c", "``").replace("\u201d", "''").replace("\u2019", "'")
    s = re.sub(r"https://(\S+?);", lambda m: "\\url{https://" + m.group(1) + "};", s)
    assert s.isascii(), s
    return s


def cover_letter(out: Path, tmp: Path):
    txt = DATE + "\n\n" + "\n\n".join(LETTER) + "\n\n" + "\n".join(SIGN) + "\n"
    (out / "05_cover_letter_JSA.txt").write_text(txt)
    body = "\n\n".join(to_latex(p) for p in LETTER)
    sign = "\\\\\n".join(to_latex(x) for x in SIGN if x and x != "Sincerely,")
    tex = rf"""\documentclass[11pt,a4paper]{{article}}
\usepackage[T1]{{fontenc}}
\usepackage{{stix}}
\usepackage[a4paper,margin=2.5cm]{{geometry}}
\usepackage{{xurl}}
\usepackage[hidelinks]{{hyperref}}
\urlstyle{{same}}
\hypersetup{{pdftitle={{Cover letter}},pdfauthor={{Anh Vu Le}}}}
\setlength{{\parindent}}{{0pt}}
\setlength{{\parskip}}{{0.7\baselineskip}}
\pagestyle{{empty}}
\begin{{document}}
\hfill {DATE}

{body}

Sincerely,

{sign}
\end{{document}}
"""
    d = tmp / "letter"
    d.mkdir()
    (d / "cover_letter.tex").write_text(tex)
    pdf = latex_build(d, "cover_letter", bib=False)
    if check_pdf(pdf).page_count != 1:
        sys.exit("cover letter is longer than one page")
    shutil.copy2(pdf, out / "05_cover_letter_JSA.pdf")


def declarations(out: Path):
    names = ", ".join(a for a, _ in AUTHORS)
    d = new_docx("Declaration of interests")
    d.add_heading("Declaration of interests", level=1)
    d.add_paragraph(f"Manuscript title: {TITLE}")
    d.add_paragraph(f"Authors: {names}")
    d.add_paragraph("\u2612 The authors declare that they have no known competing financial "
                    "interests or personal relationships that could have appeared to influence "
                    "the work reported in this paper.")
    d.add_paragraph("\u2610 The authors declare the following financial interests/personal "
                    "relationships which may be considered as potential competing interests:")
    d.save(out / "07_declaration_of_interest_JSA.docx")

    v = new_docx("Author biographies")
    v.add_heading("Author biographies (vitae)", level=1)
    v.add_paragraph("TEMPLATE: each author writes at most 100 words and supplies a "
                    "passport-type photograph as a separate image file (e.g. "
                    "Photo_<Surname>_<Given>.jpg, at least 300 dpi). Delete this note "
                    "before uploading.")
    for name, aff in AUTHORS:
        v.add_heading(name, level=2)
        v.add_paragraph(f"Affiliation: {aff}")
        v.add_paragraph("[Biography, at most 100 words: degree(s) and institution, current "
                        "position, research interests.]")
        v.add_paragraph("[Photo file name: ...]")
    v.save(out / "08_author_vitae_JSA_TEMPLATE.docx")


def readme(out: Path, main_pages: int, figs, items):
    fig_lines = "\n".join(f"  Figure_{n}.pdf  <- figures/{img}  ({lab})" for n, img, lab in figs)
    text = f"""JOURNAL OF SYSTEMS ARCHITECTURE: FILES FOR SUBMISSION ({DATE})

Manuscript: "{TITLE}"

01_JSA_main.pdf                     Main manuscript ({main_pages} pages; pdfLaTeX build of 06)
02_figures_JSA/Figure_1..{len(figs)}.pdf      Figures as separate vector files (fonts embedded)
03_highlights_JSA.docx              Highlights, {len(items)} items of at most 85 characters (Word)
03_highlights_JSA.txt               Same highlights, plain text
04_graphical_abstract_JSA.tif       Graphical abstract (TIFF, 4200 x 1620 px, 600 dpi)
05_cover_letter_JSA.pdf             Cover letter (confirm all-author approval)
05_cover_letter_JSA.txt             Cover letter, editable text
06_LaTeX_source_JSA.zip             LaTeX source (tex + bbl + bib/bst + class files +
                                    figures + generated tables)
07_declaration_of_interest_JSA.docx Declaration of interests, "nothing to declare"
                                    (same wording as the manuscript)
08_author_vitae_JSA_TEMPLATE.docx   Vitae template; NOT ready to upload

Assign each file to the corresponding portal category. Do not upload this README.
Inspect the journal-generated PDF before approving submission.

Figure files (numbering as in the manuscript):
{fig_lines}

LaTeX source: unzip 06 and compile main_jsa.tex with pdfLaTeX (pdflatex,
bibtex, pdflatex x2). The same folder, as JSA_Overleaf_Source.zip, can be
uploaded to Overleaf (New Project > Upload Project); Overleaf detects
main_jsa.tex as the main document. 01_JSA_main.pdf was produced from exactly
these files.

Before submitting:
  - Confirm all-author approval, exclusivity, and the AI declaration.
  - The Elsevier portal may ask for the declaration of interests to be produced
    with its own tool; if so, select "I have nothing to declare" and upload the
    generated .docx instead of 07 (the wording is the same).
  - Vitae: each author fills 08 (at most 100 words) and supplies a
    passport-type photograph as a separate file. These cannot be generated.
  - Data: code, raw traces and generators are at
    https://github.com/iamdinhthuan/governor-is-part-of-the-model; prediction
    files are archived on Zenodo, DOI 10.5281/zenodo.23185860. Enter the same
    data statement in the portal.
  - The graphical abstract was drawn programmatically from manuscript numbers
    (make_graphical_abstract.py); no generative image model was used.

SHA256SUMS.txt covers every file in this directory except itself.
"""
    (out / "README_FIRST.txt").write_text(text)


def sha256sums(out: Path):
    lines = []
    for p in sorted(out.rglob("*")):
        if p.is_file() and p.name != "SHA256SUMS.txt":
            lines.append(f"{hashlib.sha256(p.read_bytes()).hexdigest()}  {p.relative_to(out).as_posix()}")
    (out / "SHA256SUMS.txt").write_text("\n".join(lines) + "\n")


def main():
    for d in (PKG, OVL):
        if d.exists():
            shutil.rmtree(d)
        d.mkdir(parents=True)
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        src, build = build_source(tmp)
        pages = check_pdf(build / "main_jsa.pdf").page_count
        shutil.copy2(build / "main_jsa.pdf", PKG / "01_JSA_main.pdf")
        figs = figures(build, PKG / "02_figures_JSA")
        items = highlights(PKG)
        subprocess.run([sys.executable, str(ROOT / "make_graphical_abstract.py"), str(tmp / "ga")],
                       check=True)
        shutil.copy2(tmp / "ga" / "04_graphical_abstract_JSA.tif", PKG)
        cover_letter(PKG, tmp)
        declarations(PKG)
        write_zip(src, PKG / "06_LaTeX_source_JSA.zip")
        shutil.copytree(src, OVL / "JSA_Overleaf_Source")
        shutil.copy2(PKG / "06_LaTeX_source_JSA.zip", OVL / "JSA_Overleaf_Source.zip")
        readme(PKG, pages, figs, items)
        sha256sums(PKG)
    print(f"package: {PKG}\noverleaf: {OVL}\npages: {pages}")


if __name__ == "__main__":
    main()
